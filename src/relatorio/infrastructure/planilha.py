"""Planilha Excel dos atendimentos por médico (anexo do e-mail de atendimentos).

Uma aba por médico escolhido que teve atendimento no dia, com as colunas pedidas pela
gerência: Empresa, Funcionário, Médico, Tipo exame e Data exame. A planilha é montada em
memória no momento do envio e não é gravada em disco nem no banco (LGPD, ADR 0015).
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from relatorio.domain.exames import ExameClinico, Medico

CABECALHO = ("Empresa", "Funcionário", "Médico", "Tipo exame", "Data exame")
LARGURAS = (46, 40, 34, 22, 14)
AZUL = "164F95"
AZUL_SUAVE = "EAF1FA"
_PROIBIDOS_ABA = re.compile(r"[\[\]:*?/\\]")
MAX_NOME_ABA = 31


def _nome_aba(medico: Medico, usados: set[str]) -> str:
    base = _PROIBIDOS_ABA.sub(" ", medico.nome).strip() or medico.crm
    nome = base[:MAX_NOME_ABA].strip()
    sufixo = 2
    while nome.casefold() in usados:
        marca = f" ({sufixo})"
        nome = base[: MAX_NOME_ABA - len(marca)].strip() + marca
        sufixo += 1
    usados.add(nome.casefold())
    return nome


class GeradorPlanilhaXlsx:
    def atendimentos_por_medico(
        self, dia: date, abas: Sequence[tuple[Medico, Sequence[ExameClinico]]]
    ) -> bytes:
        livro = Workbook()
        padrao = livro.active
        if padrao is not None:
            livro.remove(padrao)
        borda = Border(bottom=Side(style="thin", color="C5CFDC"))
        usados: set[str] = set()
        for medico, exames in abas:
            aba = livro.create_sheet(_nome_aba(medico, usados))
            aba.append([f"Atendimentos de {medico.nome} (CRM {medico.crm}) em {dia:%d/%m/%Y}"])
            aba["A1"].font = Font(bold=True, size=13, color=AZUL)
            aba.append([f"{len(exames)} exame(s) clínico(s). Uso interno: contém dados pessoais."])
            aba["A2"].font = Font(italic=True, size=9, color="5F6B7A")
            aba.append(list(CABECALHO))
            for coluna in range(1, len(CABECALHO) + 1):
                celula = aba.cell(row=3, column=coluna)
                celula.font = Font(bold=True, color=AZUL)
                celula.fill = PatternFill("solid", fgColor=AZUL_SUAVE)
                celula.border = borda
            for exame in exames:
                aba.append([exame.empresa, exame.funcionario, exame.medico, exame.tipo, exame.data])
                aba.cell(row=aba.max_row, column=5).number_format = "DD/MM/YYYY"
                aba.cell(row=aba.max_row, column=5).alignment = Alignment(horizontal="center")
            for indice, largura in enumerate(LARGURAS, start=1):
                aba.column_dimensions[get_column_letter(indice)].width = largura
            aba.freeze_panes = "A4"
            aba.auto_filter.ref = f"A3:E{aba.max_row}"
            aba.print_options.gridLines = False
            aba.page_setup.orientation = "landscape"
            aba.page_setup.fitToWidth = 1
        if not livro.worksheets:  # pragma: no cover - o caso de uso só chama com atendimentos
            livro.create_sheet("Sem atendimentos")
        saida = BytesIO()
        livro.save(saida)
        return saida.getvalue()
