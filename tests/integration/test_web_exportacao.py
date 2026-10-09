"""Exportar período: pedido, barra de andamento, download e permissões (PostgreSQL)."""

from __future__ import annotations

import re
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from relatorio.domain.entidades import Agenda
from tests.fabricas import DIA
from tests.integration.test_web import (  # noqa: F401
    cliente,
    container,
    csrf,
    entrar,
    relogio,
)
from tests.integration.test_web_usuarios import criar, entrar_como, outro  # noqa: F401

pytestmark = pytest.mark.integration


def pedir(c: TestClient, token: str, tipo: str, inicio: str, fim: str) -> str:
    resposta = c.post(
        "/admin/exportacoes",
        data={"tipo": tipo, "inicio": inicio, "fim": fim},
        headers={"X-CSRF-Token": token, "HX-Request": "true"},
    )
    assert resposta.status_code == 200, resposta.text
    return resposta.text


def test_exportar_atendimento_do_pedido_ao_download(cliente: TestClient) -> None:  # noqa: F811
    token = entrar(cliente)
    pagina = cliente.get("/admin/atendimento")
    assert "Exportar período" in pagina.text and 'name="inicio"' in pagina.text
    fragmento = pedir(cliente, token, "atendimento", DIA.isoformat(), DIA.isoformat())
    assert 'role="progressbar"' in fragmento and "agua-nivel" in fragmento
    achado = re.search(r"/admin/exportacoes/([0-9a-f]{32})", fragmento)
    assert achado
    id_exportacao = achado.group(1)
    # O TestClient roda a tarefa de segundo plano ao fim da resposta: já está pronta.
    andamento = cliente.get(f"/admin/exportacoes/{id_exportacao}")
    assert "Baixar planilha" in andamento.text and 'aria-valuenow="100"' in andamento.text
    assert "hx-trigger" not in andamento.text  # parou de consultar
    arquivo = cliente.get(f"/admin/exportacoes/{id_exportacao}/arquivo")
    assert arquivo.status_code == 200
    assert "attachment" in arquivo.headers["content-disposition"]
    livro = load_workbook(BytesIO(arquivo.content))
    assert livro.sheetnames[:2] == ["Sobre", "Resumo diário"]
    # Ao reabrir a página, a última exportação aparece pronta para baixar.
    assert f"/admin/exportacoes/{id_exportacao}/arquivo" in cliente.get("/admin/atendimento").text


def test_periodo_invalido_volta_com_a_mensagem(cliente: TestClient) -> None:  # noqa: F811
    token = entrar(cliente)
    fragmento = pedir(cliente, token, "atendimento", "2026-08-01", "2026-09-15")
    assert "no máximo 31 dias" in fragmento and "progressbar" not in fragmento


def test_cada_relatorio_exige_seu_modulo(
    cliente: TestClient,  # noqa: F811
    outro: TestClient,  # noqa: F811
) -> None:
    token = entrar(cliente)
    criar(cliente, token, modulos=("atendimento",))
    assert entrar_como(outro, "ana", "senha-da-ana-1") == 303
    token_ana = csrf(outro, "/admin/atendimento")
    for tipo in ("financeiro", "sesmt", "medicos"):
        resposta = outro.post(
            "/admin/exportacoes",
            data={"tipo": tipo, "inicio": DIA.isoformat(), "fim": DIA.isoformat()},
            headers={"X-CSRF-Token": token_ana},
        )
        assert resposta.status_code == 403, tipo
    fragmento = pedir(outro, token_ana, "atendimento", DIA.isoformat(), DIA.isoformat())
    id_exportacao = re.search(r"/admin/exportacoes/([0-9a-f]{32})", fragmento)
    assert id_exportacao
    # Quem tem o módulo baixa a exportação pedida por outro usuário.
    assert cliente.get(f"/admin/exportacoes/{id_exportacao.group(1)}/arquivo").status_code == 200


def test_cartoes_nos_modulos_e_medicos_em_configuracoes(cliente: TestClient) -> None:  # noqa: F811
    entrar(cliente)
    for url, tipo in (
        ("/admin/financeiro", "financeiro"),
        ("/admin/sesmt", "sesmt"),
        ("/admin/configuracoes/medicos", "medicos"),
    ):
        pagina = cliente.get(url).text
        assert f'value="{tipo}"' in pagina and "Exportar planilha" in pagina, url
    assert "só das 20h às 5h" in cliente.get("/admin/sesmt").text
    assert 'value="medicos"' not in cliente.get("/admin/atendimento").text


def test_processar_periodo_pela_tela_sem_download(cliente: TestClient) -> None:  # noqa: F811
    token = entrar(cliente)
    pagina = cliente.get("/admin/atendimento").text
    assert "Processar período" in pagina and 'value="processar"' in pagina
    assert "Reler o dia no SGG antes" in pagina and "não reenvia e-mail" in pagina
    resposta = cliente.post(
        "/admin/exportacoes",
        data={
            "tipo": "atendimento",
            "acao": "processar",
            "inicio": DIA.isoformat(),
            "fim": DIA.isoformat(),
        },
        headers={"X-CSRF-Token": token, "HX-Request": "true"},
    )
    assert resposta.status_code == 200 and 'id="processamento-atendimento"' in resposta.text
    id_exportacao = re.search(r"/admin/exportacoes/([0-9a-f]{32})", resposta.text)
    assert id_exportacao
    andamento = cliente.get(f"/admin/exportacoes/{id_exportacao.group(1)}")
    assert "Concluído" in andamento.text and "dia(s) processado(s)" in andamento.text
    assert "Baixar planilha" not in andamento.text
    assert cliente.get(f"/admin/exportacoes/{id_exportacao.group(1)}/arquivo").status_code == 404
    # Ao reabrir a página, o último processamento aparece no seu cartão (e o de exportar não).
    reaberta = cliente.get("/admin/atendimento").text
    assert "dia(s) processado(s)" in reaberta
    assert "Baixar planilha" not in reaberta


def test_escolha_de_agendas_na_exportacao_do_atendimento(
    cliente: TestClient,  # noqa: F811
    container,  # noqa: F811
) -> None:
    token = entrar(cliente)
    with container.uow() as uow:
        uow.agendas.sincronizar(
            [
                Agenda(10, "Clínico", "Sala 1", 1, "Unidade 1"),
                Agenda(11, "Coleta", "Sala 2", 1, "Unidade 1", incluir_relatorio=False),
            ]
        )
        uow.commit()
    pagina = cliente.get("/admin/atendimento").text
    assert 'name="filtrar_agendas"' in pagina
    assert re.search(r'name="agenda" value="10" checked', pagina)  # do relatório: marcada
    assert re.search(r'name="agenda" value="11"(?! checked)', pagina)  # fora: desmarcada

    def pedir_com(dados: dict[str, object]) -> str:
        r = cliente.post(
            "/admin/exportacoes",
            data={
                "tipo": "atendimento",
                "inicio": DIA.isoformat(),
                "fim": DIA.isoformat(),
                "filtrar_agendas": "true",
                **dados,
            },
            headers={"X-CSRF-Token": token, "HX-Request": "true"},
        )
        assert r.status_code == 200, r.text
        return r.text

    assert "Marque ao menos uma agenda" in pedir_com({})
    assert "Agenda desconhecida" in pedir_com({"agenda": ["10", "999"]})
    ok = pedir_com({"agenda": ["10", "11"]})
    achado = re.search(r"/admin/exportacoes/([0-9a-f]{32})", ok)
    assert achado
    with container.uow() as uow:
        assert uow.exportacoes.obter(achado.group(1)).opcoes == "a10,a11"  # type: ignore[union-attr]
