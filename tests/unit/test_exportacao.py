"""Exportação de relatórios por período: regras, exportadores e planilha."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from relatorio.application.modelos import StatusExportacao
from relatorio.domain.entidades import FUSO_BRASILIA, AgendamentoSgg, Situacao
from relatorio.domain.exames import ExameClinico, Medico
from relatorio.domain.exportacao import Aba, Coluna, Planilha, sesmt_liberado, validar_periodo
from relatorio.domain.financeiro import ItemFaturado, Titulo
from relatorio.domain.sesmt import ContratoSesmt, DocumentoSst, EmpresaSesmt, EventoEsocial
from relatorio.infrastructure.memoria import _Exportacoes
from relatorio.infrastructure.planilha_exportacao import escrever_planilha
from tests.fabricas import DIA, hora, snapshot
from tests.unit.conftest import Sistema
from tests.unit.test_casos_de_uso import preparar_envio

HOJE = DIA + timedelta(days=1)


def exportar(s: Sistema, tipo: str, inicio: date = DIA, fim: date = DIA) -> Workbook:
    pedido = s.exportar.solicitar(tipo, inicio, fim, "admin")
    s.exportar.executar(pedido.id)
    final = s.exportar.obter(pedido.id)
    assert final is not None and final.status is StatusExportacao.PRONTO, final
    assert final.progresso == 100
    nome, conteudo = s.banco.arquivos[pedido.id]
    assert nome == f"relatorio-{tipo}-{inicio:%Y-%m-%d}-a-{fim:%Y-%m-%d}.xlsx"
    return load_workbook(BytesIO(conteudo))


def linhas(livro: Workbook, aba: str) -> list[tuple[object, ...]]:
    return [tuple(c.value for c in r) for r in livro[aba].iter_rows(min_row=4)]


# --------------------------------------------------------------------------- regras


def test_validacao_do_periodo() -> None:
    validar_periodo(DIA, DIA + timedelta(days=30), DIA + timedelta(days=40))
    with pytest.raises(ValueError, match="posterior"):
        validar_periodo(DIA, DIA - timedelta(days=1), HOJE)
    with pytest.raises(ValueError, match="futuro"):
        validar_periodo(DIA, HOJE + timedelta(days=1), HOJE)
    with pytest.raises(ValueError, match="31 dias"):
        validar_periodo(DIA - timedelta(days=31), DIA, HOJE)


@pytest.mark.parametrize(
    ("hora_", "liberado"), [(19, False), (20, True), (23, True), (0, True), (4, True), (5, False)]
)
def test_janela_do_sesmt(hora_: int, liberado: bool) -> None:
    assert sesmt_liberado(datetime(2026, 10, 7, hora_, 30, tzinfo=FUSO_BRASILIA)) is liberado


def test_sesmt_fora_do_horario_nem_entra_na_fila(sistema: Sistema) -> None:
    sistema.relogio.instante = hora("10:00", HOJE)
    with pytest.raises(ValueError, match="20h às 5h"):
        sistema.exportar.solicitar("sesmt", DIA, DIA, "admin")
    assert sistema.banco.exportacoes == {}


def test_tipo_desconhecido(sistema: Sistema) -> None:
    with pytest.raises(ValueError, match="desconhecido"):
        sistema.exportar.solicitar("rh", DIA, DIA, "admin")


def test_planilha_tem_capa_relatorio_antes_da_fonte_e_formatos() -> None:
    planilha = Planilha("teste", "Relatório de teste", DIA, DIA, observacoes=["Uma observação."])
    planilha.abas = [
        Aba("Fonte - X", "fonte", (Coluna("Valor", "moeda"),), [(12.5,)], fonte=True),
        Aba(
            "Resumo diário",
            "relatório",
            (Coluna("Data", "data"), Coluna("Tempo", "duracao"), Coluna("Taxa", "percentual")),
            [(DIA, 125, 0.25)],
        ),
        Aba("Vazia", "sem linhas", (Coluna("A"),)),
    ]
    livro = load_workbook(BytesIO(escrever_planilha(planilha, hora("10:00"))))
    assert livro.sheetnames == ["Sobre", "Resumo diário", "Vazia", "Fonte - X"]
    capa = [tuple(c.value for c in r) for r in livro["Sobre"].iter_rows()]
    assert ("Período", "23/09/2026 a 23/09/2026") in capa
    assert ("Observação", "Uma observação.") in capa
    resumo = livro["Resumo diário"]
    assert [c.value for c in resumo[3]] == ["Data", "Tempo", "Taxa"]
    assert (
        resumo["B4"].value == timedelta(seconds=125) and resumo["B4"].number_format == "[h]:mm:ss"
    )
    assert resumo["C4"].number_format == "0.0%"
    assert resumo.freeze_panes == "A4"
    assert livro["Fonte - X"]["A4"].number_format.startswith('"R$"')
    assert livro["Vazia"]["A4"].value == "Sem dados no período."


# --------------------------------------------------------------------------- atendimento


def test_atendimento_exporta_resumo_e_fonte_do_banco(sistema: Sistema) -> None:
    preparar_envio(sistema)
    livro = exportar(sistema, "atendimento")
    assert livro.sheetnames == [
        "Sobre",
        "Resumo diário",
        "Por agenda",
        "Fonte - Agendamentos",
        "Fonte - Mudanças de status",
    ]
    (dia,) = linhas(livro, "Resumo diário")
    assert dia[0].date() == DIA and dia[3] == 1  # um atendimento
    (agendamento,) = linhas(livro, "Fonte - Agendamentos")
    assert (agendamento[0], agendamento[3], agendamento[9]) == (1, "Clínico", "Atendido")
    assert [r[4] for r in linhas(livro, "Fonte - Mudanças de status")] == [
        "Aguardando",
        "Em Atendimento",
        "Atendido",
    ]


def test_atendimento_traz_o_tipo_de_exame_lido_do_sgg(sistema: Sistema) -> None:
    preparar_envio(sistema)
    sistema.sgg.do_dia = [
        AgendamentoSgg(
            1, "Clínico", 1, "Unidade 1", DIA, time(8, 0), Situacao.ATENDIDO, None, tipo="Periódico"
        )
    ]
    livro = exportar(sistema, "atendimento")
    cabecalho = [c.value for c in livro["Fonte - Agendamentos"][3]]
    assert cabecalho[4] == "Tipo de exame"
    (agendamento,) = linhas(livro, "Fonte - Agendamentos")
    assert agendamento[4] == "Periódico"


def test_atendimento_sem_resposta_do_sgg_deixa_o_tipo_em_branco_e_avisa(
    sistema: Sistema,
) -> None:
    preparar_envio(sistema)
    sistema.sgg.falhar = True
    livro = exportar(sistema, "atendimento")
    (agendamento,) = linhas(livro, "Fonte - Agendamentos")
    assert agendamento[4] in (None, "")
    capa = [r[1] for r in livro["Sobre"].iter_rows(values_only=True)]
    assert any("tipo de exame ficou em branco" in str(v) for v in capa)


def test_atendimento_ordena_agendamentos_com_e_sem_horario(sistema: Sistema) -> None:
    preparar_envio(sistema)
    uow = sistema.uow()
    uow.snapshots.salvar(snapshot(2, Situacao.AGENDADO, hora_agendamento=None))
    uow.snapshots.salvar(snapshot(3, Situacao.AGENDADO, hora_agendamento=time(7, 0)))
    livro = exportar(sistema, "atendimento")
    ids = [r[0] for r in linhas(livro, "Fonte - Agendamentos")]
    assert ids[0] == 2 and set(ids) == {1, 2, 3}  # sem horário vem primeiro


def test_progresso_e_gravado_e_o_dia_sem_resumo_e_avisado(
    sistema: Sistema, monkeypatch: pytest.MonkeyPatch
) -> None:
    preparar_envio(sistema)
    etapas: list[tuple[int, str]] = []
    original = _Exportacoes.progresso

    def espiar(self: _Exportacoes, id_exportacao: str, progresso: int, etapa: str) -> None:
        etapas.append((progresso, etapa))
        original(self, id_exportacao, progresso, etapa)

    monkeypatch.setattr(_Exportacoes, "progresso", espiar)
    pedido = sistema.exportar.solicitar("atendimento", DIA - timedelta(days=1), DIA, "admin")
    sistema.exportar.executar(pedido.id)
    progressos = [p for p, _ in etapas]
    assert progressos == sorted(progressos) and progressos[0] == 1 and progressos[-1] == 97
    assert any("agendamentos" in e for _, e in etapas)
    _, conteudo = sistema.banco.arquivos[pedido.id]
    capa = [r[1] for r in load_workbook(BytesIO(conteudo))["Sobre"].iter_rows(values_only=True)]
    assert any(isinstance(v, str) and "1 dia(s) sem resumo: 22/09" in v for v in capa)


# --------------------------------------------------------------------------- financeiro


def titulo(
    id_: int, valor: str, pagamento: date | None = None, emissao: date | None = None
) -> Titulo:
    return Titulo(
        id_,
        Decimal(valor),
        "Pago",
        False,
        emissao or DIA,
        DIA,
        pagamento,
        nome=f"Cliente {id_} LTDA",
        itens=(ItemFaturado(1, "PCMSO", 1, Decimal(valor)),),
    )


def test_financeiro_rele_os_titulos_no_sgg(sistema: Sistema) -> None:
    f = sistema.financeiro
    f.recebidos = [titulo(1, "150.00", DIA), titulo(2, "99.90", DIA - timedelta(days=40))]
    f.pagos = [titulo(3, "80.00", DIA)]
    f.emitidos = [titulo(4, "300.00", emissao=DIA)]
    livro = exportar(sistema, "financeiro")
    assert [r[0] for r in linhas(livro, "Fonte - Recebimentos")] == [1]  # o de fora do período não
    assert linhas(livro, "Fonte - Recebimentos")[0][6] == 150.0
    assert [r[1] for r in linhas(livro, "Fonte - Pagamentos")] == ["Cliente 3 LTDA"]
    assert linhas(livro, "Fonte - Itens faturados")[0][3:] == ("PCMSO", 1, 300.0)
    assert {"receber_pagos", "pagar_pagos", "receber_emitidos"} <= set(f.chamadas)


def test_sgg_fora_marca_falha_com_mensagem(sistema: Sistema) -> None:
    sistema.financeiro.falhar = True
    pedido = sistema.exportar.solicitar("financeiro", DIA, DIA, "admin")
    with pytest.raises(Exception, match="indisponível"):
        sistema.exportar.executar(pedido.id)
    final = sistema.exportar.obter(pedido.id)
    assert final is not None and final.status is StatusExportacao.FALHA
    assert final.erro.startswith("O SGG não respondeu")
    assert sistema.exportar.executar(pedido.id)["status"] == "ignorada"  # não roda de novo


# --------------------------------------------------------------------------- SESMT


def test_sesmt_rele_o_sgg_e_filtra_eventos_do_periodo(sistema: Sistema) -> None:
    s = sistema.sesmt
    s.empresas = [EmpresaSesmt(1, "Conplan LTDA", "", "1", True)]
    s.contratos = [ContratoSesmt(10, 1, date(2026, 10, 30), "Em andamento", True, "2025-1")]
    s.documentos = {1: [DocumentoSst(1, "PGR", date(2026, 11, 1))]}
    s.eventos = {
        1: [
            EventoEsocial("a", "S-2220", 1, 50, DIA, None, None, "1.1", ""),
            EventoEsocial("b", "S-2240", 1, 51, DIA - timedelta(days=5), None, None, "", ""),
        ]
    }
    sistema.relogio.instante = hora("21:00", HOJE)
    livro = exportar(sistema, "sesmt")
    assert [r[1] for r in linhas(livro, "Fonte - Empresas")] == ["Conplan LTDA"]
    assert [r[2] for r in linhas(livro, "Fonte - Contratos")] == ["Conplan LTDA"]
    assert [r[1] for r in linhas(livro, "Fonte - Documentos")] == ["PGR"]
    eventos = linhas(livro, "Fonte - Eventos eSocial")
    assert [(r[0], r[3]) for r in eventos] == [("a", "Funcionário #50")]  # só o do período


# --------------------------------------------------------------------------- médicos


def test_medicos_traz_nomes_so_na_planilha(sistema: Sistema) -> None:
    ana = Medico("34985-DF", "Ana Souza")
    sistema.uow().medicos.adicionar(ana)
    sistema.exames.exames[DIA] = [
        ExameClinico(1, DIA, 10, 501, ana.crm, ana.nome, "Admissional", funcionario="Fulano"),
        ExameClinico(2, DIA, 10, 502, ana.crm, ana.nome, "Periódico", funcionario="Beltrano"),
    ]
    sistema.exames.empresas = {10: EmpresaSesmt(10, "Conplan LTDA", "", "1", False)}
    sistema.processar_exames.executar(DIA)
    livro = exportar(sistema, "medicos", DIA - timedelta(days=1), DIA)
    (dia,) = linhas(livro, "Por dia e médico")
    assert dia[2:5] == ("Ana Souza", "34985-DF", 2) and dia[5:7] == (1, 1)
    assert sorted(r[2] for r in linhas(livro, "Fonte - Exames clínicos")) == ["Beltrano", "Fulano"]
    assert all(e.funcionario == "" for e in sistema.banco.exames.values())  # banco sem nomes
    capa = [r[1] for r in livro["Sobre"].iter_rows(values_only=True)]
    assert any(isinstance(v, str) and "sem processamento dos exames: 22/09" in v for v in capa)


def test_medicos_sem_sgg_saem_pelo_codigo(sistema: Sistema) -> None:
    ana = Medico("34985-DF", "Ana Souza")
    sistema.uow().medicos.adicionar(ana)
    sistema.exames.exames[DIA] = [ExameClinico(1, DIA, 10, 501, ana.crm, ana.nome, "Outro")]
    sistema.processar_exames.executar(DIA)
    sistema.exames.falhar = True
    livro = exportar(sistema, "medicos")
    assert linhas(livro, "Fonte - Exames clínicos")[0][2] == "Funcionário #501"


def test_ultimas_por_usuario_e_limpeza_de_24h(sistema: Sistema) -> None:
    preparar_envio(sistema)
    antiga = sistema.exportar.solicitar("atendimento", DIA, DIA, "ana")
    sistema.relogio.instante += timedelta(hours=25)
    nova = sistema.exportar.solicitar("atendimento", DIA, DIA, "admin")
    assert antiga.id not in sistema.banco.exportacoes  # apagada ao pedir a nova
    assert sistema.exportar.ultimas("admin") == {"exportar:atendimento": nova}
    assert sistema.exportar.ultimas("ana") == {}


# --------------------------------------------------------------------------- processar período


def processar(s: Sistema, tipo: str, inicio: date = DIA, fim: date = DIA, opcoes: str = "") -> str:
    pedido = s.exportar.solicitar(tipo, inicio, fim, "admin", "processar", opcoes)
    resultado = s.exportar.executar(pedido.id)
    final = s.exportar.obter(pedido.id)
    assert final is not None and final.status is StatusExportacao.PRONTO and final.progresso == 100
    assert pedido.id not in s.banco.arquivos  # processamento não gera arquivo
    assert resultado["resultado"] == final.etapa
    return final.etapa


def test_processar_periodo_do_atendimento_recalcula_cada_dia_sem_enviar(sistema: Sistema) -> None:
    preparar_envio(sistema)
    sistema.banco.resumos.clear()
    sistema.relogio.instante = hora("10:00", HOJE + timedelta(days=1))
    etapa = processar(sistema, "atendimento", DIA - timedelta(days=1), DIA)
    assert etapa == "2 dia(s) processado(s)"
    assert DIA in sistema.banco.resumos and (DIA - timedelta(days=1)) in sistema.banco.resumos
    assert sistema.email.enviados == []  # nunca reenvia e-mail no período


def test_processar_atendimento_com_releitura_no_sgg_so_se_pedido(sistema: Sistema) -> None:
    preparar_envio(sistema)
    sistema.relogio.instante = hora("10:00", HOJE + timedelta(days=1))
    antes = sistema.sgg.requisicoes_realizadas
    processar(sistema, "atendimento")
    sem_releitura = sistema.sgg.requisicoes_realizadas
    processar(sistema, "atendimento", opcoes="reconciliar")
    assert sistema.sgg.requisicoes_realizadas > sem_releitura >= antes


def test_processar_financeiro_le_o_sgg_de_cada_dia(sistema: Sistema) -> None:
    sistema.relogio.instante = hora("10:00", HOJE + timedelta(days=2))
    etapa = processar(sistema, "financeiro", DIA - timedelta(days=1), DIA)
    assert etapa == "2 dia(s) processado(s)"
    assert {DIA, DIA - timedelta(days=1)} <= set(sistema.banco.resumos_financeiros)


def test_processar_financeiro_tolera_dia_sem_sgg_mas_nao_todos(sistema: Sistema) -> None:
    sistema.relogio.instante = hora("10:00", HOJE + timedelta(days=2))
    sistema.financeiro.falhar = True
    pedido = sistema.exportar.solicitar(
        "financeiro", DIA - timedelta(days=1), DIA, "admin", "processar"
    )
    with pytest.raises(Exception, match="nenhum dia"):
        sistema.exportar.executar(pedido.id)
    final = sistema.exportar.obter(pedido.id)
    assert final is not None and final.status is StatusExportacao.FALHA


def test_processar_sesmt_uma_coleta_para_varios_dias(sistema: Sistema) -> None:
    s = sistema.sesmt
    s.empresas = [EmpresaSesmt(1, "Conplan LTDA", "", "1", True)]
    s.contratos = [ContratoSesmt(10, 1, date(2026, 10, 30), "Em andamento", True, "2025-1")]
    s.eventos = {
        1: [
            EventoEsocial("a", "S-2220", 1, 50, DIA, None, None, "1.1", ""),
            EventoEsocial("b", "S-2240", 1, 51, DIA - timedelta(days=1), None, None, "", ""),
        ]
    }
    sistema.relogio.instante = hora("22:00", HOJE)
    etapa = processar(sistema, "sesmt", DIA - timedelta(days=1), DIA)
    assert etapa == "2 dia(s) processado(s)"
    assert s.chamadas.count("empresas") == 1  # uma leitura só
    totais = {d: r.metricas["esocial"]["total"] for d, r in sistema.banco.resumos_sesmt.items()}
    assert totais == {DIA: 1, DIA - timedelta(days=1): 1}
    assert {d: r.metricas["referencia"] for d, r in sistema.banco.resumos_sesmt.items()} == {
        DIA: DIA.isoformat(),
        DIA - timedelta(days=1): (DIA - timedelta(days=1)).isoformat(),
    }


def test_processar_sesmt_so_na_janela_e_so_ate_ontem(sistema: Sistema) -> None:
    sistema.relogio.instante = hora("10:00", HOJE)
    with pytest.raises(ValueError, match=r"O processamento do SESMT.*20h às 5h"):
        sistema.exportar.solicitar("sesmt", DIA, DIA, "admin", "processar")
    sistema.relogio.instante = hora("22:00", HOJE)
    with pytest.raises(ValueError, match="dia anterior"):
        sistema.exportar.solicitar("sesmt", DIA, HOJE, "admin", "processar")
    with pytest.raises(ValueError, match="dia anterior"):
        sistema.exportar.solicitar("financeiro", DIA, HOJE, "admin", "processar")
    sistema.exportar.solicitar("atendimento", DIA, HOJE, "admin", "processar")  # hoje vale


def test_processar_medicos_le_os_exames_de_cada_dia(sistema: Sistema) -> None:
    ana = Medico("34985-DF", "Ana Souza")
    sistema.uow().medicos.adicionar(ana)
    antes = DIA - timedelta(days=1)
    sistema.exames.exames = {
        DIA: [ExameClinico(1, DIA, 10, 501, ana.crm, ana.nome, "Outro", funcionario="X")],
        antes: [ExameClinico(2, antes, 10, 502, ana.crm, ana.nome, "Periódico", funcionario="Y")],
    }
    sistema.exames.empresas = {10: EmpresaSesmt(10, "Conplan LTDA", "", "1", False)}
    etapa = processar(sistema, "medicos", antes, DIA)
    assert etapa == "2 dia(s) processado(s)"
    assert set(sistema.banco.exames) == {1, 2}


def test_acao_desconhecida(sistema: Sistema) -> None:
    with pytest.raises(ValueError, match="desconhecido"):
        sistema.exportar.solicitar("atendimento", DIA, DIA, "admin", "apagar")


def test_ultimas_separa_exportar_e_processar(sistema: Sistema) -> None:
    preparar_envio(sistema)
    a = sistema.exportar.solicitar("atendimento", DIA, DIA, "admin")
    b = sistema.exportar.solicitar("atendimento", DIA, DIA, "admin", "processar")
    assert sistema.exportar.ultimas("admin") == {
        "exportar:atendimento": a,
        "processar:atendimento": b,
    }
