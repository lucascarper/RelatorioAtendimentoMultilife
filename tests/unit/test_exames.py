"""Atendimentos por médico: regras, processamento noturno, e-mail e planilha anexa."""

from __future__ import annotations

from datetime import date, timedelta
from email import message_from_bytes
from io import BytesIO

import pytest
from openpyxl import load_workbook

from relatorio.application.exames import TIPO_XLSX
from relatorio.application.modelos import Anexo, ConteudoEmail
from relatorio.domain.exames import (
    ExameClinico,
    Medico,
    linhas_da_planilha,
    normalizar_crm,
    resumo_por_medico,
)
from relatorio.domain.sesmt import EmpresaSesmt
from relatorio.infrastructure.email.envio import montar_mensagem
from relatorio.infrastructure.planilha import GeradorPlanilhaXlsx
from tests.fabricas import DIA, hora
from tests.unit.conftest import Sistema
from tests.unit.test_casos_de_uso import preparar_envio

ANA = Medico("34985-DF", "Ana Souza")
BETO = Medico("31096-DF", "Beto Lima")
CAIO = Medico("85686-MG", "Caio Reis")


def exame(
    id_: int, medico: Medico, tipo: str = "Admissional", empresa: int = 1, funcionario: str = ""
) -> ExameClinico:
    return ExameClinico(
        id=id_,
        data=DIA,
        id_empresa=empresa,
        id_funcionario=1000 + id_,
        crm=medico.crm,
        medico=medico.nome,
        tipo=tipo,
        funcionario=funcionario or f"Trabalhador {id_}",
    )


def autorizar_planilha(sistema: Sistema, *ids: int) -> None:
    """Marca destinatários (de ``preparar_envio``: 1 = glauco, 2 = tecnologia)."""
    uow = sistema.uow()
    for id_destinatario in ids:
        uow.destinatarios.definir_anexo(id_destinatario, True)


def escolher(sistema: Sistema, *medicos: Medico) -> None:
    uow = sistema.uow()
    for medico in medicos:
        uow.medicos.adicionar(medico)


@pytest.fixture
def com_exames(sistema: Sistema) -> Sistema:
    sistema.exames.exames[DIA] = [
        exame(1, ANA, "Admissional", empresa=10),
        exame(2, ANA, "Periódico", empresa=10),
        exame(3, ANA, "Admissional", empresa=11),
        exame(4, BETO, "Demissional", empresa=11),
        exame(5, CAIO, "Outro", empresa=12),  # médico não escolhido
    ]
    sistema.exames.empresas = {
        10: EmpresaSesmt(10, "Conplan Sistemas LTDA", "", "1", False),
        11: EmpresaSesmt(11, "Metalúrgica Sul LTDA", "", "1", False),
        12: EmpresaSesmt(12, "Posto Central LTDA", "", "1", False),
    }
    return sistema


# --------------------------------------------------------------------------- regras


@pytest.mark.parametrize(
    ("texto", "esperado"),
    [("34985-DF", "34985-DF"), ("34985df", "34985-DF"), (" 5201174436 - rj ", "5201174436-RJ")],
)
def test_crm_normalizado(texto: str, esperado: str) -> None:
    assert normalizar_crm(texto) == esperado


@pytest.mark.parametrize("texto", ["", "DF", "1-DF", "34985-D", "abc-DF"])
def test_crm_invalido(texto: str) -> None:
    with pytest.raises(ValueError, match="CRM"):
        normalizar_crm(texto)


def test_resumo_inclui_medico_sem_atendimento_e_ordena_por_total() -> None:
    exames = [exame(1, ANA, "Periódico"), exame(2, ANA, "Admissional"), exame(3, BETO)]
    resumo = resumo_por_medico(exames, [BETO, ANA, CAIO], DIA)
    assert resumo["total"] == 3
    assert [(m["nome"], m["total"]) for m in resumo["medicos"]] == [
        ("Ana Souza", 2),
        ("Beto Lima", 1),
        ("Caio Reis", 0),
    ]
    # Tipos na ordem do cadastro do SGG, não na ordem de chegada.
    assert [t["tipo"] for t in resumo["medicos"][0]["por_tipo"]] == ["Admissional", "Periódico"]


def test_planilha_tem_uma_aba_por_medico_com_atendimento() -> None:
    exames = [exame(1, BETO), exame(2, ANA), exame(3, ANA)]
    abas = linhas_da_planilha(exames, [ANA, BETO, CAIO])
    assert [(m.nome, len(linhas)) for m, linhas in abas] == [("Ana Souza", 2), ("Beto Lima", 1)]


def test_planilha_xlsx_colunas_e_nomes_de_aba() -> None:
    longo = Medico("1-DF", "Dra. Maria/Clara: Nome Muito Comprido Para Uma Aba")
    abas = [
        (ANA, [exame(1, ANA, funcionario="Fulano")]),
        (longo, [exame(2, longo)]),
        (Medico("2-DF", longo.nome), [exame(3, longo)]),
    ]
    conteudo = GeradorPlanilhaXlsx().atendimentos_por_medico(DIA, abas)
    livro = load_workbook(BytesIO(conteudo))
    assert livro.sheetnames[0] == "Ana Souza"
    assert all(len(nome) <= 31 and "/" not in nome for nome in livro.sheetnames)
    assert len(set(livro.sheetnames)) == 3
    aba = livro["Ana Souza"]
    assert [c.value for c in aba[3]] == [
        "Empresa",
        "Funcionário",
        "Médico",
        "Tipo exame",
        "Data exame",
    ]
    linha = [c.value for c in aba[4]]
    assert linha[1:4] == ["Fulano", "Ana Souza", "Admissional"]
    assert linha[4].date() == DIA


# --------------------------------------------------------------------------- processamento


def test_processa_registra_medicos_e_grava_so_os_escolhidos_sem_nome(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    detalhe = s.processar_exames.executar(DIA)
    assert detalhe == {
        "referencia": DIA.isoformat(),
        "clinicos": 5,
        "selecionados": 3,
        "medicos_vistos": 3,
        "medicos_escolhidos": 1,
        "empresas_consultadas": 2,
        "empresas_sem_nome": 0,
    }
    assert {m.crm for m in s.banco.medicos.values()} == {ANA.crm, BETO.crm, CAIO.crm}
    assert not s.banco.medicos[BETO.crm].selecionado
    assert s.banco.medicos[CAIO.crm].visto_em == DIA
    gravados = s.banco.exames.values()
    assert {e.id for e in gravados} == {1, 2, 3}
    assert all(e.funcionario == "" for e in gravados)  # LGPD: o nome nunca é gravado
    assert {e.empresa for e in gravados} == {"Conplan Sistemas LTDA", "Metalúrgica Sul LTDA"}


def test_empresa_ja_conhecida_nao_e_consultada_de_novo(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    s.processar_exames.executar(DIA)
    amanha = DIA + timedelta(days=1)
    s.exames.exames[amanha] = [
        ExameClinico(9, amanha, 10, 7, ANA.crm, ANA.nome, "Periódico", funcionario="X")
    ]
    s.exames.chamadas.clear()
    s.processar_exames.executar(amanha)
    assert s.exames.chamadas == [f"exames:{amanha.isoformat()}"]
    assert s.banco.exames[9].empresa == "Conplan Sistemas LTDA"


def test_empresa_que_falha_aparece_pelo_codigo(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    s.exames.falhar_empresas = {11}
    detalhe = s.processar_exames.executar(DIA)
    assert detalhe["empresas_sem_nome"] == 1
    assert s.banco.exames[3].empresa == "Empresa #11"


def test_reprocessar_substitui_o_dia(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA, BETO)
    s.processar_exames.executar(DIA)
    s.exames.exames[DIA] = [exame(4, BETO, "Demissional", empresa=11)]
    s.processar_exames.executar(DIA)
    assert set(s.banco.exames) == {4}
    assert s.banco.coletas_exames[DIA].clinicos == 1


def test_terceira_falha_do_processamento_alerta(com_exames: Sistema) -> None:
    s = com_exames
    s.exames.falhar = True
    for tentativa in (1, 2):
        with pytest.raises(Exception, match="indisponível"):
            s.processar_exames.executar_agendado(DIA, tentativa)
    assert s.alertas_enviados() == []
    with pytest.raises(Exception, match="indisponível"):
        s.processar_exames.executar_agendado(DIA, 3)
    assert s.alertas_enviados() == ["[Alerta] Atendimentos por médico não processados"]


# --------------------------------------------------------------------------- e-mail


def test_sem_medico_escolhido_o_email_sai_como_antes(com_exames: Sistema) -> None:
    s = com_exames
    preparar_envio(s)
    assert s.enviar.executar(DIA)["status"] == "enviado"
    _, conteudo = s.email.enviados[0]
    assert conteudo.anexos == ()
    assert conteudo.texto == "texto"
    assert s.exames.chamadas == []  # nenhuma consulta ao SGG à toa


def test_email_traz_resumo_e_planilha_com_nomes(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA, BETO)
    s.processar_exames.executar(DIA)
    preparar_envio(s)
    autorizar_planilha(s, 1, 2)
    resultado = s.enviar.executar(DIA)
    assert resultado["anexos"] == 1
    _, conteudo = s.email.enviados[0]
    assert conteudo.texto == "exames=4"
    (anexo,) = conteudo.anexos
    assert anexo.nome == f"atendimentos-por-medico-{DIA:%Y-%m-%d}.xlsx"
    assert anexo.tipo == TIPO_XLSX
    livro = load_workbook(BytesIO(anexo.conteudo))
    assert livro.sheetnames == ["Ana Souza", "Beto Lima"]
    nomes = sorted(str(linha[1].value) for linha in livro["Ana Souza"].iter_rows(min_row=4))
    assert nomes == ["Trabalhador 1", "Trabalhador 2", "Trabalhador 3"]
    # Os nomes só passaram pela memória: o banco continua sem eles.
    assert all(e.funcionario == "" for e in s.banco.exames.values())


def test_envio_processa_quando_a_noite_nao_rodou_ou_a_escolha_mudou(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    s.processar_exames.executar(DIA)
    escolher(s, BETO)  # escolha mudou depois do processamento
    preparar_envio(s)
    autorizar_planilha(s, 1, 2)
    s.enviar.executar(DIA)
    _, conteudo = s.email.enviados[0]
    assert conteudo.texto == "exames=4"
    assert s.banco.coletas_exames[DIA].medicos == tuple(sorted((ANA.crm, BETO.crm)))


def test_sem_sgg_no_envio_a_planilha_sai_com_codigo(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    s.processar_exames.executar(DIA)
    preparar_envio(s)
    autorizar_planilha(s, 1, 2)
    s.exames.falhar = True
    s.enviar.executar(DIA)
    (anexo,) = s.email.enviados[0][1].anexos
    aba = load_workbook(BytesIO(anexo.conteudo))["Ana Souza"]
    assert sorted(str(c[1].value) for c in aba.iter_rows(min_row=4)) == [
        "Funcionário #1001",
        "Funcionário #1002",
        "Funcionário #1003",
    ]


def test_complemento_que_falha_nao_segura_o_relatorio(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    preparar_envio(s)
    autorizar_planilha(s, 1, 2)
    s.exames.falhar = True  # nem processado à noite, nem no envio
    assert s.enviar.executar(DIA)["status"] == "enviado"
    relatorio = [c for _, c in s.email.enviados if not c.assunto.startswith("[Alerta]")]
    assert relatorio[0].anexos == ()
    assert "[Alerta] Relatório enviado sem os atendimentos por médico" in s.alertas_enviados()


def test_dia_sem_exame_dos_escolhidos_tem_secao_sem_anexo(com_exames: Sistema) -> None:
    s = com_exames
    s.exames.exames[DIA] = [exame(5, CAIO, "Outro", empresa=12)]
    escolher(s, ANA)
    preparar_envio(s)
    autorizar_planilha(s, 1, 2)
    s.enviar.executar(DIA)
    _, conteudo = s.email.enviados[0]
    assert conteudo.texto == "exames=0"
    assert conteudo.anexos == ()


def test_previa_do_admin_le_so_o_banco(com_exames: Sistema) -> None:
    s = com_exames
    assert s.complemento_exames.resumo_gravado(DIA) == {}
    escolher(s, ANA)
    assert s.complemento_exames.resumo_gravado(DIA) == {}  # ainda não processado
    s.processar_exames.executar(DIA)
    s.exames.chamadas.clear()
    resumo = s.complemento_exames.resumo_gravado(DIA)["exames_medicos"]
    assert resumo["total"] == 3
    assert s.exames.chamadas == []


def test_planilha_so_para_quem_esta_autorizado(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    s.processar_exames.executar(DIA)
    preparar_envio(s)
    autorizar_planilha(s, 1)
    resultado = s.enviar.executar(DIA)
    assert (resultado["destinatarios"], resultado["com_anexo"]) == (2, 1)
    (com, conteudo_com), (sem, conteudo_sem) = s.email.enviados
    assert com == ["glauco@multilife.com.br"] and len(conteudo_com.anexos) == 1
    assert conteudo_com.texto == "exames=3"
    assert sem == ["tecnologia@multilife.com.br"] and conteudo_sem.anexos == ()
    assert conteudo_sem.texto == "exames=3 restrito"


def test_ninguem_autorizado_nem_le_nomes_no_sgg(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    s.processar_exames.executar(DIA)
    preparar_envio(s)
    s.exames.chamadas.clear()
    s.enviar.executar(DIA)
    ((destinatarios, conteudo),) = s.email.enviados
    assert len(destinatarios) == 2 and conteudo.anexos == ()
    assert conteudo.texto == "exames=3 restrito"
    assert s.exames.chamadas == []  # sem planilha, os nomes não são lidos


def test_retencao_apaga_exames_antigos(com_exames: Sistema) -> None:
    s = com_exames
    escolher(s, ANA)
    s.processar_exames.executar(DIA)
    s.relogio.instante = hora("03:00", date(DIA.year + 3, 1, 1))
    assert s.limpar.executar()["exames"] == 3
    assert not s.banco.exames and not s.banco.coletas_exames


def test_mensagem_com_anexo_vira_mixed() -> None:
    conteudo = ConteudoEmail(
        assunto="Resumo",
        html="<p>oi</p>",
        texto="oi",
        anexos=(Anexo("planilha.xlsx", b"PK\x03\x04", TIPO_XLSX),),
    )
    mensagem = message_from_bytes(bytes(montar_mensagem("a@b.com", "A", ["c@d.com"], conteudo)))
    assert mensagem.get_content_type() == "multipart/mixed"
    partes = [p for p in mensagem.walk() if p.get_filename() == "planilha.xlsx"]
    assert len(partes) == 1
    assert partes[0].get_content_type() == TIPO_XLSX
    assert partes[0].get_payload(decode=True) == b"PK\x03\x04"
